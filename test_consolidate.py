import json
from pathlib import Path

from openpyxl import load_workbook

import make_sample_data
from consolidate import consolidate

CFG = json.loads(Path(__file__).with_name("config.json").read_text(encoding="utf-8"))


def test_incremental_and_scoped_to_module(tmp_path):
    src = tmp_path / "exports"
    make_sample_data.main(src, orders=2, quotations=1)
    # Two modules where only the second one has a width: the first must not inherit it.
    (src / "export_5999999_Test Customer.csv").write_text(
        "5999999;BoQ export\n;;MOD-A;3;;;;;01.02.2025\n;;;;LEN_MM;1200\n;;MOD-B;1;;;;;\n;;;;WIDTH_MM;620\n",
        encoding="latin1",
    )
    (src / "notes.csv").write_text("not an export\n", encoding="latin1")
    book = tmp_path / "out" / "db.xlsx"

    first = consolidate([src], book, CFG)
    assert first["Sales Orders"] > 0 and first["Quotations"] > 0
    assert consolidate([src], book, CFG) == {"Sales Orders": 0, "Quotations": 0}

    ws = load_workbook(book)["Sales Orders"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True) if str(r[0]) == "5999999"]
    assert [r["Module"] for r in rows] == ["MOD-A", "MOD-B"]
    assert rows[0]["Quantity"] == 3 and rows[0]["Length (mm)"] == "1200" and rows[0]["Width (mm)"] is None
    assert rows[1]["Width (mm)"] == "620" and rows[0]["Delivery Date"] == "01/02/2025"
    assert rows[0]["Customer"] == "Test Customer"
    assert ws.tables["SalesOrdersTable"].ref == f"A1:N{ws.max_row}"
