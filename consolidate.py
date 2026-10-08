"""Consolidate bill-of-quantities CSV exports into one Excel workbook for BI.

Each export is one sales order or quotation. The first cell holds the document
number; module lines carry a product code and quantity; the lines below a module
carry its characteristics as key/value pairs. See README.md for the layout.
"""
import argparse
import csv
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FIXED_HEAD = ["Document", "Module", "Quantity"]
FIXED_TAIL = ["Creation Date", "Delivery Date", "Customer"]


def parse_export(path, cfg):
    """Return (sheet_name, document_number, rows) or None if the file is not an export."""
    with open(path, newline="", encoding=cfg["csv"]["encoding"]) as f:
        lines = list(csv.reader(f, delimiter=cfg["csv"]["delimiter"]))
    if not lines or not lines[0]:
        return None
    number = lines[0][0].strip()
    sheet = next((s for prefix, s in cfg["document_types"].items() if number.startswith(prefix)), None)
    if sheet is None:
        return None

    key_to_column = {key: column for column, key in cfg["columns"].items()}
    shared = {
        "Document": number,
        "Creation Date": datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d"),
        "Customer": path.stem.rsplit("_", 1)[-1].strip(),
    }
    rows, current = [], None
    for line in lines[1:]:
        code = line[2].strip() if len(line) > 2 else ""
        if code.startswith(cfg["module_prefix"]):
            quantity = line[3].strip() if len(line) > 3 else ""
            current = dict(shared, Module=code, Quantity=int(quantity) if quantity.isdigit() else None)
            current["Delivery Date"] = line[8].strip().replace(".", "/") if len(line) > 8 else ""
            rows.append(current)
        elif current is not None:
            # A characteristic belongs to the module above it, so the search stops at the next module line.
            for i, cell in enumerate(line[:-1]):
                column = key_to_column.get(cell.strip())
                if column and column not in current:
                    current[column] = line[i + 1].strip()
    return sheet, number, rows


def consolidate(source_dirs, workbook_path, cfg):
    """Append exports not yet in the workbook. Returns {sheet_name: rows_added}."""
    header = FIXED_HEAD + list(cfg["columns"]) + FIXED_TAIL
    sheets = list(cfg["document_types"].values())
    workbook_path = Path(workbook_path)
    if workbook_path.exists():
        wb = load_workbook(workbook_path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    for name in sheets:
        if name not in wb.sheetnames:
            wb.create_sheet(name).append(header)

    known = {name: {str(row[0]) for row in wb[name].iter_rows(min_row=2, max_col=1, values_only=True)} for name in sheets}
    added = dict.fromkeys(sheets, 0)
    for directory in source_dirs:
        for path in sorted(Path(directory).glob("*.csv")):
            try:
                parsed = parse_export(path, cfg)
            except (OSError, UnicodeDecodeError, csv.Error) as exc:
                print(f"skipped {path.name}: {exc}", file=sys.stderr)
                continue
            if parsed is None:
                continue
            sheet, number, rows = parsed
            if number in known[sheet]:
                continue
            known[sheet].add(number)
            for row in rows:
                wb[sheet].append([row.get(column) for column in header])
            added[sheet] += len(rows)

    for name in sheets:
        ws = wb[name]
        if ws.max_row < 2:
            continue  # Excel rejects a table with a header and no data rows
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        if ws.tables:
            next(iter(ws.tables.values())).ref = ref
        else:
            table = Table(displayName=name.replace(" ", "") + "Table", ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
            ws.add_table(table)
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(workbook_path)
    return added


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("source", nargs="+", help="folder(s) containing the CSV exports")
    parser.add_argument("-o", "--output", default="output/database.xlsx", help="workbook to create or update")
    parser.add_argument("-c", "--config", default=Path(__file__).with_name("config.json"))
    args = parser.parse_args()
    cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
    for sheet, count in consolidate(args.source, args.output, cfg).items():
        print(f"{sheet}: {count} rows added")


if __name__ == "__main__":
    main()
